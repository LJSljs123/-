"""把史料录入 Excel 转成 events.json。

用法：
    python src/import_excel.py
    python src/import_excel.py data/curated/史料录入模板.xlsx -o data/staging/events.from_xlsx.json

规则：
- 同一 id 的多行会合并成一条事件，多行视为多条来源。
- 必填缺失、枚举非法、原文摘录为空的行会被拒绝，并逐条报告。
- 有错误的导入只允许写到 data/staging/，正式库不会被部分结果覆盖。
- 此旧表不含完整主库字段，禁止直接覆盖 data/curated/events.json。
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import sys
from collections import OrderedDict
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from safe_write import atomic_write_text

ROOT = Path(__file__).resolve().parent.parent

EVENT_TYPES = {"战斗", "转移", "密营建立", "被捕", "牺牲", "会议", "群众工作", "其他"}
DATE_PREC = {"日": "day", "月": "month", "年": "year"}
PLACE_PREC = {"遗址": "site", "乡镇": "town", "县": "county", "省": "province"}
REL = {"A", "B", "C"}


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_date(raw: str, prec_cn: str) -> tuple[str | None, str, str | None]:
    """返回 (start, precision, problem)。"""
    s = raw.replace("/", "-").replace(".", "-").strip()
    m = re.match(r"^(\d{4})(?:-(\d{1,2}))?(?:-(\d{1,2}))?$", s)
    if not m:
        return None, "", f"日期格式无法解析：{raw}"
    y, mo, d = m.group(1), m.group(2), m.group(3)
    prec = DATE_PREC.get(prec_cn)
    if prec_cn and prec is None:
        return None, "", f"日期精度无效：{prec_cn}"
    if not prec:
        prec = "day" if d else "month" if mo else "year"
    try:
        # Validate every component present in the sheet before reducing it to
        # month/year precision. date() catches month/day overflow and year 0000.
        date(int(y), int(mo or 1), int(d or 1))
    except ValueError:
        return None, "", f"日期不是有效日历日期：{raw}"
    if prec == "day":
        if not (mo and d):
            return None, "", f"日期精度标为“日”但缺日：{raw}"
        start = f"{y}-{int(mo):02d}-{int(d):02d}"
    elif prec == "month":
        if not mo:
            return None, "", f"日期精度标为“月”但缺月：{raw}"
        start = f"{y}-{int(mo):02d}-01"
    else:
        start = f"{y}-01-01"
    return start, prec, None


def num(raw: str) -> float | None:
    try:
        value = float(raw) if raw else None
        return round(value, 5) if value is not None and math.isfinite(value) else None
    except ValueError:
        return None


def int_if(v: float | None) -> float | int | None:
    """45000.0 -> 45000，避免 JSON 里出现无意义的小数。"""
    if v is None:
        return None
    return int(v) if float(v).is_integer() else v


def read_measure(r: list, idx: int, eid: str, book: str, page: str,
                 problems: list[str]) -> dict | None:
    """从一行里读数量项。三列全空视为没有；部分填写视为录入错误。"""
    name, raw_val, unit = cell(r[17]), cell(r[18]), cell(r[19])
    if not (name or raw_val or unit):
        return None
    if not (name and raw_val and unit):
        problems.append(f"第{idx}行（{eid}）：数量项填写不完整（名称/数值/单位缺一）")
        return None
    val = num(raw_val)
    if val is None:
        problems.append(f"第{idx}行（{eid}）：数量项数值不是数字（{raw_val}）")
        return None
    return {"name": name, "value": int_if(val), "unit": unit,
            "source": book, "page": page or None}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", type=Path, nargs="?", default=ROOT / "data" / "curated" / "史料录入模板.xlsx")
    ap.add_argument("-o", "--out", type=Path,
                    default=ROOT / "data" / "staging" / "events.from_xlsx.json",
                    help="默认写 data/staging/，不覆盖正式库")
    ap.add_argument("--force", action="store_true",
                    help="允许写入 data/curated/ 的非主库文件（会保留原文件备份）；events.json 始终禁止覆盖")
    args = ap.parse_args()

    if not args.xlsx.exists():
        print(f"找不到 {args.xlsx}", file=sys.stderr)
        return 1

    ws = load_workbook(args.xlsx, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(min_row=3, values_only=True))

    problems: list[str] = []
    merged: "OrderedDict[str, dict]" = OrderedDict()
    seen_source: dict[str, set] = {}

    for idx, r in enumerate(rows, start=3):
        if not r or not any(cell(x) for x in r):
            continue
        r = list(r) + [None] * (20 - len(r))
        eid = cell(r[0])
        if not eid:
            problems.append(f"第{idx}行：缺 id")
            continue

        src_book, src_page, src_quote, src_rel = cell(r[12]), cell(r[13]), cell(r[14]), cell(r[15])
        if not (src_book and src_quote and src_rel in REL):
            problems.append(f"第{idx}行（{eid}）：来源不完整（书名/原文/可信度任一缺失）")
            continue

        if eid not in merged:
            start, prec, prob = parse_date(cell(r[1]), cell(r[2]))
            if prob:
                problems.append(f"第{idx}行（{eid}）：{prob}")
                continue
            lat_raw, lng_raw = cell(r[5]), cell(r[6])
            lat, lng = num(lat_raw), num(lng_raw)
            if (lat_raw and lat is None) or (lng_raw and lng is None):
                problems.append(f"第{idx}行（{eid}）：经纬度必须是有限数字")
            if lat is not None and not -90 <= lat <= 90:
                problems.append(f"第{idx}行（{eid}）：纬度超出 -90 到 90 的范围")
            if lng is not None and not -180 <= lng <= 180:
                problems.append(f"第{idx}行（{eid}）：经度超出 -180 到 180 的范围")
            etype = cell(r[10])
            if etype not in EVENT_TYPES:
                problems.append(f"第{idx}行（{eid}）：事件类型非法（{etype}）")
                continue
            pprec_cn = cell(r[7])
            if pprec_cn not in PLACE_PREC:
                problems.append(f"第{idx}行（{eid}）：地点精度非法（{pprec_cn}）")
                continue
            summary = cell(r[11])
            if not summary:
                problems.append(f"第{idx}行（{eid}）：摘要为空")
                continue

            units = [x for x in re.split(r"[、,，]", cell(r[8])) if x]
            persons = [x for x in re.split(r"[、,，]", cell(r[9])) if x]

            merged[eid] = {
                "id": eid,
                "date": {"start": start, "end": None, "precision": prec},
                "place": {
                    "name": cell(r[3]),
                    "modern": cell(r[4]) or cell(r[3]),
                    "lat": lat,
                    "lng": lng,
                    "precision": PLACE_PREC[pprec_cn],
                },
                "units": units,
                "persons": persons,
                "event_type": etype,
                "summary": summary,
                "casualties": {"our": None, "enemy": None},
                "measures": [],
                "sources": [],
                "confidence": 0,
                "conflict_group": cell(r[16]) or None,
            }
            seen_source[eid] = set()

        key = (src_book, src_page, src_quote)
        if key in seen_source[eid]:
            continue
        seen_source[eid].add(key)
        merged[eid]["sources"].append(
            {
                "book": src_book,
                "page": src_page or None,
                "quote": src_quote,
                "reliability": src_rel,
            }
        )
        m = read_measure(r, idx, eid, src_book, src_page, problems)
        if m:
            merged[eid]["measures"].append(m)

    events = list(merged.values())

    # P0-10 / P1-11：正式库只接受明确授权且无坏行的完整导入。
    out_resolved = args.out.resolve()
    curated_dir = (ROOT / "data" / "curated").resolve()
    staging_dir = (ROOT / "data" / "staging").resolve()
    canonical_events = (curated_dir / "events.json").resolve()
    writes_curated = curated_dir in out_resolved.parents
    writes_staging = staging_dir == out_resolved or staging_dir in out_resolved.parents
    if out_resolved == canonical_events:
        print("拒绝覆盖 data/curated/events.json：录入表是部分来源，不含完整主库字段。")
        print("请先输出到 data/staging/，人工合并并保留 date_variants、provenance 等字段。")
        return 2
    if curated_dir in out_resolved.parents:
        if not args.force:
            print("拒绝写入 data/curated/：该目录是正式库，须显式 --force。")
            print("默认请写到 data/staging/，人工比对后再替换。")
            return 2
    if not events:
        problems.append("录入表没有可用事件，拒绝输出空库")
    if problems and (writes_curated or not writes_staging):
        print(f"读入 {len(rows)} 行，生成 {len(events)} 条事件，问题 {len(problems)} 条")
        for problem in problems[:30]:
            print("  !", problem)
        if len(problems) > 30:
            print(f"  ... 另有 {len(problems) - 30} 条")
        print("存在错误：正式目录和非 staging 路径均未写入；完整问题导出请显式指定 data/staging/。")
        return 1

    args.out.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(
        {
            "_note": "由 src/import_excel.py 从录入表生成，请勿手工编辑。",
            "events": events,
        },
        ensure_ascii=False,
        indent=1,
    )
    if writes_curated and args.out.exists():
        backup = args.out.with_suffix(args.out.suffix + ".bak")
        index = 1
        while backup.exists():
            backup = args.out.with_suffix(args.out.suffix + f".{index}.bak")
            index += 1
        shutil.copy2(args.out, backup)
        print(f"已备份原文件 → {backup}")
    atomic_write_text(args.out, payload)

    print(f"读入 {len(rows)} 行，生成 {len(events)} 条事件，问题 {len(problems)} 条")
    for p in problems[:30]:
        print("  !", p)
    if len(problems) > 30:
        print(f"  ... 另有 {len(problems) - 30} 条")
    print(f"已写出 {args.out}")
    if args.out.resolve().parent.name == "staging":
        print("注意：这是 staging 副本，正式库未被改动。人工比对后再替换 data/curated/events.json。")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
