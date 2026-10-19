"""生成史料录入 Excel 模板。

用法：python src/make_template.py
产出：data/curated/史料录入模板.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "curated" / "史料录入模板.xlsx"

COLS = [
    ("id", 22, "唯一标识，如 evt_1940_jingyu_01。同一事件多条来源时，多行写相同 id"),
    ("日期", 13, "YYYY-MM-DD；只知道月份写 YYYY-MM；只知道年份写 YYYY"),
    ("日期精度", 10, "日 / 月 / 年"),
    ("历史地名", 20, "史料中的写法"),
    ("今地名", 20, "现行行政区划名称"),
    ("纬度", 10, "GCJ-02；只知到县就填县质心"),
    ("经度", 10, "GCJ-02"),
    ("地点精度", 10, "遗址 / 乡镇 / 县 / 省"),
    ("部队番号", 22, "按史料原文，多个用顿号分隔"),
    ("相关人物", 22, "多个用顿号分隔"),
    ("事件类型", 12, "战斗 / 转移 / 密营建立 / 被捕 / 牺牲 / 会议 / 群众工作 / 其他"),
    ("摘要", 46, "不超过 400 字，客观陈述，不加评论"),
    ("来源书名", 30, "书名或馆藏名"),
    ("页码", 10, "页码 / 馆藏编号 / 档案号"),
    ("原文摘录", 52, "照抄原文，不得改写。这是最重要的一列"),
    ("可信度", 10, "A=一手档案 B=正式出版物/权威二手报道 C=回忆录/口述/展陈/待核转述"),
    ("冲突组", 18, "同一事件的多条不同记载填相同编号，如 cg_1940_jingyu"),
    ("数量项", 14, "该条来源给出的数字叫什么，如“总兵力”“歼敌数”“作战次数”。没有就留空"),
    ("数值", 10, "纯数字。原文写“四万五千余”填 45000，写“3万余人”填 30000"),
    ("单位", 8, "人 / 次 / 支 / 件"),
]

REQUIRED = {0, 1, 2, 3, 4, 7, 10, 11, 12, 14, 15}


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "史料录入"

    head_fill = PatternFill("solid", fgColor="1F2937")
    req_fill = PatternFill("solid", fgColor="7F1D1D")
    body_font = Font(name="微软雅黑", size=10)
    head_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")

    for i, (name, width, note) in enumerate(COLS, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = req_fill if (i - 1) in REQUIRED else head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = width
        d = ws.cell(row=2, column=i, value=note)
        d.font = Font(name="微软雅黑", size=9, color="6B7280")
        d.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 40

    ws.freeze_panes = "A3"

    dv_date = DataValidation(type="list", formula1='"日,月,年"', allow_blank=True)
    dv_place = DataValidation(type="list", formula1='"遗址,乡镇,县,省"', allow_blank=True)
    dv_type = DataValidation(
        type="list",
        formula1='"战斗,转移,密营建立,被捕,牺牲,会议,群众工作,其他"',
        allow_blank=True,
    )
    dv_rel = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
    for dv, col in ((dv_date, "C"), (dv_place, "H"), (dv_type, "K"), (dv_rel, "P")):
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}1000")

    for row in range(3, 1001):
        for col in range(1, len(COLS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = body_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    guide = wb.create_sheet("填写须知")
    rules = [
        ("填写须知", ""),
        ("", ""),
        ("红色表头为必填项。", ""),
        ("1", "原文摘录必须照抄，不得改写、不得概括。这一列是作品的可信度来源，也是评委会看的地方。"),
        ("2", "不知道的字段留空，不要猜。日期只到月就写 YYYY-MM，并把日期精度选为“月”。"),
        ("3", "同一事件如果有两条以上不同记载，写多行、id 相同、冲突组填同一个编号。差异由脚本自动比对。"),
        ("4", "同一事件的多条来源也写多行、id 相同，脚本会自动合并为一条事件的多条来源。"),
        ("5", "可信度：A=一手档案（原件、档案号）；B=正式出版物或权威二手报道；C=回忆录、口述、展陈说明或待核转述。"),
        ("6", "录入完成后运行：python src/import_excel.py，脚本会校验并生成 data/staging/events.from_xlsx.json。"),
        ("7", "有错误时脚本非零退出；草稿仅供审阅，不能直接替换 data/curated/events.json。"),
        ("8", "「数量项 / 数值 / 单位」三列用来记录史料里的数字（兵力、歼敌数、作战次数）。"
              "数字必须跟着来源走：同一事件的两条来源给出不同数字时，分别写在各自那一行。"
              "脚本会自动发现同一条事件内部的数值分歧，不需要额外标注冲突组。"),
        ("9", "只写季节不写月份时（如“1942年秋”），日期按该季节首月填写，精度选“月”，"
              "并在原文摘录里保留“秋”的原话，不要改写。"),
    ]
    for i, (a, b) in enumerate(rules, start=1):
        guide.cell(row=i, column=1, value=a).font = Font(
            name="微软雅黑", size=11, bold=(i <= 2)
        )
        guide.cell(row=i, column=2, value=b).font = Font(name="微软雅黑", size=10)
        guide.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 8
    guide.column_dimensions["B"].width = 110
    for i in range(3, len(rules) + 1):
        guide.row_dimensions[i].height = 44

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"已生成 {OUT}")


if __name__ == "__main__":
    main()
